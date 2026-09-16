"""导出「试剂验收」待办清单 Excel，供逐条录入 5 个样本的比对结果。

用法：python scripts/report_lot_verifications.py
"""
import json
import urllib.parse
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

H = "http://lab-management-282724-9-1408547492.sh.run.tcloudbase.com"
OUT = "outputs/试剂验收待办清单.xlsx"
HEAD_FILL = PatternFill("solid", fgColor="1A365D")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")


def login():
    data = urllib.parse.urlencode(
        {"username": "jinzizheng", "password": "Jzz6827556"}).encode()
    r = urllib.request.Request(H + "/api/v1/auth/login", data=data)
    return json.load(urllib.request.urlopen(r, timeout=60))["access_token"]


def get(tok, path):
    r = urllib.request.Request(H + path,
                               headers={"Authorization": "Bearer " + tok})
    return json.load(urllib.request.urlopen(r, timeout=300))


def main():
    tok = login()
    rows, page = [], 1
    while True:
        d = get(tok, f"/api/v1/reagent/lot-verifications?page={page}&page_size=200")
        rows += d["items"]
        if len(rows) >= d["total"] or not d["items"]:
            break
        page += 1

    wb = Workbook()
    ws = wb.active
    ws.title = "待验收清单"
    heads = ["序号", "责任库", "类型", "试剂 / 质控品", "规格", "品牌",
             "旧批号", "旧批号效期", "新批号", "新批号效期", "变更日期",
             "检验项目", "允许偏倚%", "判定标准", "结论", "状态", "备注"]
    ws.append(heads)
    for i, v in enumerate(rows, 1):
        ws.append([
            i, v.get("library", ""), v.get("item_type", ""), v.get("reagent_name", ""),
            v.get("spec", ""), v.get("brand", ""),
            v.get("old_batch_no", ""), str(v.get("old_expiry_date") or ""),
            v.get("new_batch_no", ""), str(v.get("new_expiry_date") or ""),
            str(v.get("change_date") or ""),
            v.get("test_item_name", ""), v.get("allow_bias_pct", ""),
            v.get("criterion_label", ""), v.get("conclusion", ""),
            v.get("status", ""), v.get("remark", ""),
        ])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, w in enumerate([6, 10, 9, 34, 20, 12, 15, 13, 15, 13, 12, 18, 11, 34, 11, 10, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="center", wrap_text=True)
        if not str(row[12].value or "").strip():
            for c in row:
                c.fill = WARN_FILL        # 缺允许偏倚 → 橙色
        elif str(row[14].value) == "符合要求":
            for c in row:
                c.fill = OK_FILL

    ws2 = wb.create_sheet("汇总")
    ws2.append(["项目", "数量"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    miss = sum(1 for v in rows if not str(v.get("allow_bias_pct") or "").strip())
    for k, n in [("验收记录总数", len(rows)),
                 ("待验证", sum(1 for v in rows if v.get("status") == "待验证")),
                 ("已完成", sum(1 for v in rows if v.get("status") == "已完成")),
                 ("未匹配到判定标准（需手工填允许偏倚）", miss),
                 ("免疫", sum(1 for v in rows if v.get("library") == "免疫")),
                 ("生化凝血", sum(1 for v in rows if v.get("library") == "生化凝血")),
                 ("试剂", sum(1 for v in rows if v.get("item_type") == "试剂")),
                 ("质控品", sum(1 for v in rows if v.get("item_type") == "质控品"))]:
        ws2.append([k, n])
    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["B"].width = 10
    ws2.append([])
    ws2.append(["填表说明", "① 橙色行=未自动匹配到允许偏倚，请手工填写；"
                          "② 判定规则：5 个样本中 ≥4 个相对偏倚绝对值 ≤ 允许偏倚 → 符合要求；"
                          "③ 样本可为质控品+患者样本混合"])

    wb.save(OUT)
    print(f"已生成 {OUT}：共 {len(rows)} 条，其中 {miss} 条需手工填允许偏倚")
    print(f"  待验证 {sum(1 for v in rows if v.get('status') == '待验证')} 条，"
          f"已完成 {sum(1 for v in rows if v.get('status') == '已完成')} 条")


if __name__ == "__main__":
    main()
