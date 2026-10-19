# -*- coding: utf-8 -*-
"""组内培训：自动生成「培训考题及成绩汇总」「培训效果评价汇总」Word 文档。

数据来源：
- TrainingSession（培训基本信息）
- PreJobAuth（关联本培训的考核单：岗位、仪器、权限）
- ExamBank（岗位题库：理论题含答案、实操要点、口头问答）
- PreJobTheoryRecord（答题成绩、提交时间）
- PreJobAuth.exam_json["_satisfaction"]（培训满意度 1~5）
"""
from __future__ import annotations

import json
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PASS_LINE = 60          # 理论合格线
EXCELLENT_LINE = 90     # 优秀线
SATISFY_LINE = 4        # 满意度 ≥4 视为满意

_CN_NUM = "一二三四五六七八九十"


def _load(v, default):
    if isinstance(v, (list, dict)):
        return v
    try:
        return json.loads(v or "")
    except Exception:
        return default


def _set_cn_font(doc: Document):
    st = doc.styles["Normal"]
    st.font.name = "宋体"
    st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(16)
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def _h(doc: Document, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    p.paragraph_format.space_before = Pt(8)


def _table(doc: Document, rows: list[list[str]], widths: list[float] | None = None, header=True):
    if not rows:
        return
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 and header) or j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            r.font.size = Pt(9)
            if i == 0 and header:
                r.bold = True
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t


def _info_rows(sess) -> list[list[str]]:
    return [
        ["培训时间", sess.train_time or "—", "培训人", sess.teacher or "—"],
        ["培训地点", sess.location or "—", "培训对象", sess.target or "—"],
        ["培训内容", sess.content or "—", "培训类别", sess.tag or "—"],
    ]


def _collect(db, sess, PreJobAuth, PreJobTheoryRecord):
    """汇总本培训关联的考核单、答题记录、岗位题库。"""
    from ..models.education import ExamBank  # 延迟导入，避免循环

    auths = (db.query(PreJobAuth)
               .filter(PreJobAuth.training_session_id == sess.id)
               .order_by(PreJobAuth.name)
               .all())
    rows = []
    posts = []
    for a in auths:
        recs = (db.query(PreJobTheoryRecord)
                  .filter(PreJobTheoryRecord.pre_job_auth_id == a.id)
                  .order_by(PreJobTheoryRecord.submit_at)
                  .all())
        ej = _load(a.exam_json, {})
        sat = ej.get("_satisfaction") if isinstance(ej, dict) else None
        best = max((r.score_pct or 0) for r in recs) if recs else None
        last = recs[-1] if recs else None
        apos = _load(a.positions_json, [])
        for p in apos:
            if p not in posts:
                posts.append(p)
        rows.append({
            "name": a.name,
            "positions": apos,
            "instruments": _load(a.instruments_json, []),
            "permissions": _load(a.permissions_json, []),
            "best": best,
            "attempts": len(recs),
            "satisfaction": sat,
            "submit_at": last.submit_at.strftime("%Y-%m-%d %H:%M") if last and last.submit_at else "—",
        })
    banks = {}
    for p in posts:
        b = db.query(ExamBank).filter(ExamBank.post == p).first()
        if not b:
            continue
        banks[p] = {
            "theory": _load(b.theory_json, {}),
            "practical": _load(b.practical_json, []),
            "qa": _load(b.qa_json, []),
        }
    return rows, posts, banks


def build_exam_summary(db, sess, PreJobAuth, PreJobTheoryRecord) -> bytes:
    """生成「培训考题及成绩汇总」docx。"""
    rows, posts, banks = _collect(db, sess, PreJobAuth, PreJobTheoryRecord)

    doc = Document()
    _set_cn_font(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)

    _title(doc, f"{sess.name or '培训'} — 培训考题及成绩汇总")
    _table(doc, _info_rows(sess), widths=[2.6, 5.2, 2.6, 5.2])

    # 一、理论考题
    _h(doc, "一、理论考题（含标准答案）")
    if not posts:
        doc.add_paragraph("（本次培训未关联岗位题库）")
    for pi, post in enumerate(posts):
        bank = banks.get(post)
        if not bank:
            continue
        th = bank["theory"]
        p = doc.add_paragraph()
        r = p.add_run(f"（{_CN_NUM[pi] if pi < 10 else pi + 1}）{post}")
        r.bold = True
        n = 0
        for key, label in (("single", "单选题"), ("multi", "多选题"), ("judge", "判断题")):
            qs = th.get(key) or []
            if not qs:
                continue
            doc.add_paragraph(f"{label}（{len(qs)} 题）").runs[0].bold = True
            for q in qs:
                n += 1
                doc.add_paragraph(f"{n}. {q.get('q', '')}")
                opts = q.get("options") or []
                if opts:
                    doc.add_paragraph("　　" + "　".join(str(o) for o in opts))
                ans = q.get("answer")
                if isinstance(ans, (list, tuple)):
                    ans = "".join(str(x) for x in ans)
                doc.add_paragraph(f"　　标准答案：{ans}")

    # 二、实操考核要点
    _h(doc, "二、实操考核要点（满分 100 分，合格线 80 分）")
    has_practical = False
    for post in posts:
        bank = banks.get(post) or {}
        items = bank.get("practical") or []
        if not items:
            continue
        has_practical = True
        doc.add_paragraph(f"{post}：").runs[0].bold = True
        data = [["序号", "考核要点", "分值"]]
        for i, it in enumerate(items, 1):
            data.append([str(i), it.get("point", ""), str(it.get("score", ""))])
        data.append(["", "合计", str(sum(int(it.get("score") or 0) for it in items))])
        _table(doc, data, widths=[1.4, 12.6, 2.0])
    if not has_practical:
        doc.add_paragraph("（未配置实操要点）")

    # 三、口头问答
    _h(doc, "三、口头问答题目")
    has_qa = False
    for post in posts:
        bank = banks.get(post) or {}
        qs = bank.get("qa") or []
        if not qs:
            continue
        has_qa = True
        doc.add_paragraph(f"{post}：").runs[0].bold = True
        for i, q in enumerate(qs, 1):
            doc.add_paragraph(f"{i}. {q.get('q', '')}")
    if not has_qa:
        doc.add_paragraph("（未配置口头问答题）")

    # 四、成绩汇总
    _h(doc, "四、成绩汇总")
    data = [["序号", "姓名", "考核岗位", "理论得分", "是否合格", "满意度", "提交时间"]]
    total = passed = 0
    for i, r in enumerate(rows, 1):
        best = r["best"]
        ok = "—" if best is None else ("合格" if best >= PASS_LINE else "不合格")
        if best is not None:
            total += 1
            if best >= PASS_LINE:
                passed += 1
        sat = r["satisfaction"]
        sat_txt = f"{sat} 分" if sat else "—"
        data.append([
            str(i), r["name"], "、".join(r["positions"]) or "—",
            f"{best}" if best is not None else "未考", ok, sat_txt, r["submit_at"],
        ])
    _table(doc, data, widths=[1.2, 2.2, 3.4, 2.2, 2.2, 2.0, 3.4])

    # 五、结论
    _h(doc, "五、结论")
    rate = round(passed * 100.0 / total, 1) if total else 0
    doc.add_paragraph(
        f"本次培训应参加考核 {len(rows)} 人，实际参加理论考核 {total} 人，"
        f"合格 {passed} 人，合格率 {rate}%。"
        + ("理论考核全部通过。" if total and passed == total else "")
        + ("仍有人员未参加或未通过理论考核，需补考后再行授权。" if total < len(rows) or passed < total else "")
    )
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_eval_summary(db, sess, PreJobAuth, PreJobTheoryRecord) -> bytes:
    """生成「培训效果评价汇总」docx。"""
    rows, posts, _banks = _collect(db, sess, PreJobAuth, PreJobTheoryRecord)

    doc = Document()
    _set_cn_font(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)

    _title(doc, f"{sess.name or '培训'} — 培训效果评价汇总")
    _table(doc, _info_rows(sess), widths=[2.6, 5.2, 2.6, 5.2])

    _h(doc, "一、评价方式")
    doc.add_paragraph(
        "受训人员在完成扫码理论考核后，同步对本次培训的收获程度进行评价"
        "（5 分制：5 非常满意、4 满意、3 一般、2 较差、1 很差）。"
        "满意率 = 评分 ≥4 分的人数 / 参评人数。"
    )

    _h(doc, "二、评价明细")
    data = [["序号", "姓名", "岗位", "评分（5 分制）", "评价等级", "评价时间"]]
    vals = []
    for i, r in enumerate(rows, 1):
        v = r["satisfaction"]
        if v:
            vals.append(float(v))
        grade = {5: "非常满意", 4: "满意", 3: "一般", 2: "较差", 1: "很差"}.get(int(v) if v else 0, "未评价")
        data.append([str(i), r["name"], "、".join(r["positions"]) or "—",
                     f"{v}" if v else "—", grade, r["submit_at"]])
    _table(doc, data, widths=[1.2, 2.4, 3.8, 3.0, 3.0, 4.0])

    _h(doc, "三、统计汇总")
    n = len(vals)
    satisfy = sum(1 for v in vals if v >= SATISFY_LINE)
    avg = round(sum(vals) / n, 2) if n else 0
    rate = round(satisfy * 100.0 / n, 1) if n else 0
    _table(doc, [
        ["参评人数", f"{n} 人", "满意人数（≥4 分）", f"{satisfy} 人"],
        ["满意率", f"{rate}%", "平均分", f"{avg} 分"],
    ], widths=[3.4, 3.6, 4.2, 3.4], header=False)

    _h(doc, "四、结论")
    if n:
        level = "很好" if rate >= 90 else ("良好" if rate >= 80 else "一般，建议改进培训方式")
        doc.add_paragraph(
            f"本次培训共 {n} 人参与效果评价，满意率 {rate}%，平均 {avg} 分，培训效果{level}。"
        )
    else:
        doc.add_paragraph("暂无人员完成效果评价，建议督促受训人员扫码答题时同步完成评价。")
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
