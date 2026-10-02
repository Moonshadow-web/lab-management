# -*- coding: utf-8 -*-
"""设备卡片（仪器标识卡 + 维修二维码）生成服务

输出：Word (.docx)
版式：A4 竖版，每页 3 张卡片；每张卡片 = 1 个 3 列 × 10 行表格
     卡片总宽 9.21cm（标签 3.50 / 值 3.30 / 二维码 2.41），可裁剪

要素来源（仪器档案）：
    设备编号 dept_no / 设备名称 name / 厂家型号 model / 设备状态 status
    设备负责人 owner / 开始使用日期 start_date / 本次校准、下次校准（CalibrationRecord）
    设备维修联系方式 repair_contact（默认 3000，允许修改）
"""
import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt
from docx.text.paragraph import Paragraph

WML = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
# 列宽 twips（3.50 / 3.30 / 2.41 cm = 9.21cm 总宽）
COL_W = [1985, 1871, 1364]
CARD_W = sum(COL_W)
QR_PX_CM = 1.60          # 二维码图片宽度(cm)
PER_PAGE = 3             # 每页卡片数
BLANK_LINES = 3          # 卡片之间的空行数


def ym(v: str) -> str:
    """把各种日期写法归一成「YYYY 年 M 月」（卡片只显示到年月，不显示具体日子）"""
    if not v:
        return ''
    s = str(v).strip()
    m = re.search(r'(\d{4})\s*[-/年.]\s*(\d{1,2})', s)
    if m:
        return f'{m.group(1)} 年 {int(m.group(2))} 月'
    m = re.search(r'(\d{4})', s)
    return f'{m.group(1)} 年' if m else s


def _set_cell_text(cell, text, size=9, bold=False):
    """覆盖单元格文本（保留首段格式）"""
    paras = cell._tc.findall(qn('w:p'))
    if not paras:
        return
    p = paras[0]
    for ch in list(p):
        if ch.tag != qn('w:pPr'):
            p.remove(ch)
    for ex in paras[1:]:
        ex.getparent().remove(ex)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rPr.append(sz)
    if bold:
        b = OxmlElement('w:b'); rPr.append(b)
    r.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t)
    p.append(r)


def _tight_para(cell, txt, size=5.5):
    """在单元格末尾追加一个紧凑居中段落"""
    xml = (f'<w:p {WML}><w:pPr>'
           f'<w:spacing w:before="0" w:after="0" w:line="200" w:lineRule="exact"/>'
           f'<w:jc w:val="center"/>'
           f'<w:rPr><w:sz w:val="{int(size*2)}"/><w:szCs w:val="{int(size*2)}"/></w:rPr>'
           f'</w:pPr>'
           f'<w:r><w:rPr><w:sz w:val="{int(size*2)}"/><w:szCs w:val="{int(size*2)}"/></w:rPr>'
           f'<w:t xml:space="preserve">{txt}</w:t></w:r></w:p>')
    cell._tc.append(parse_xml(xml))


def _make_qr_png(url: str, path: str):
    """生成二维码 PNG；无 qrcode 库时返回 None（卡片仍可生成，只是无图）"""
    try:
        import qrcode
        from qrcode.constants import ERROR_CORRECT_M
    except Exception:
        return None
    if os.path.exists(path):
        return path
    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_M, box_size=24, border=1)
    qr.add_data(url)
    qr.make(fit=True)
    qr.make_image(fill_color='black', back_color='white').convert('RGB').save(path)
    return path


def _shade(cell, fill='FFFFFF'):
    tcPr = cell._tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr'); cell._tc.insert(0, tcPr)
    for tag in ('w:shd', 'w:tcShd'):
        e = tcPr.find(qn(tag))
        if e is not None:
            tcPr.remove(e)
    shd = OxmlElement('w:tcShd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _valign(cell, val='center'):
    tcPr = cell._tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr'); cell._tc.insert(0, tcPr)
    e = tcPr.find(qn('w:vAlign'))
    if e is not None:
        tcPr.remove(e)
    va = OxmlElement('w:vAlign'); va.set(qn('w:val'), val); tcPr.append(va)


def build_card_table(doc, item, qr_dir, host):
    """在 doc 末尾追加一张设备卡片表格，返回表格对象"""
    code = (item.get('dept_no') or '').strip()
    name = (item.get('name') or '').strip()
    model = (item.get('model') or '').strip()
    status = (item.get('status') or '在用').strip()
    owner = (item.get('owner') or '').strip()
    start = ym(item.get('start_date'))
    cal = ym(item.get('cal_date'))
    nextcal = ym(item.get('next_cal_date'))
    contact = (item.get('repair_contact') or '3000').strip()

    tbl = doc.add_table(rows=10, cols=3)
    tbl.style = 'Table Grid'
    tbl.autofit = False
    # 列宽（三处同步）
    grid = tbl._tbl.find(qn('w:tblGrid'))
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(COL_W[i]))
    for tr in tbl._tbl.findall(qn('w:tr')):
        for i, tc in enumerate(tr.findall(qn('w:tc'))):
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
            o = tcPr.find(qn('w:tcW'))
            if o is not None:
                tcPr.remove(o)
            w = OxmlElement('w:tcW'); w.set(qn('w:w'), str(COL_W[i])); w.set(qn('w:type'), 'dxa')
            tcPr.append(w)
    tblW = tbl._tbl.find(qn('w:tblPr')).find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tbl._tbl.find(qn('w:tblPr')).append(tblW)
    tblW.set(qn('w:w'), str(CARD_W)); tblW.set(qn('w:type'), 'dxa')

    # 标题行跨 3 列
    tbl.rows[0].cells[0].merge(tbl.rows[0].cells[2])
    # R7-R10 值列跨 2-3 列
    for ri in range(6, 10):
        tbl.rows[ri].cells[1].merge(tbl.rows[ri].cells[2])
    # 二维码列 R2-R6 纵向合并
    cq = tbl.rows[1].cells[2]
    for i in range(2, 6):
        cq = cq.merge(tbl.rows[i].cells[2])

    rows_data = [
        ('设备编号', code),
        ('设备名称', name),
        ('厂家型号', model),
        ('设备状态', status),
        ('设备负责人', owner),
        ('开始使用日期', start),
        ('本次校准时间', cal),
        ('下次校准时间', nextcal),
        ('设备维修联系方式', contact),
    ]
    _set_cell_text(tbl.rows[0].cells[0], '民航总医院检验科设备卡片', 11, True)
    for i, (k, v) in enumerate(rows_data, start=1):
        _set_cell_text(tbl.rows[i].cells[0], k, 9)
        _set_cell_text(tbl.rows[i].cells[1], v, 9 if k != '设备状态' else 7.5)

    # 二维码单元格
    cell = tbl.rows[1].cells[2]
    _shade(cell, 'FFFFFF')
    _valign(cell, 'center')
    tc = cell._tc
    ps = tc.findall(qn('w:p'))
    for p in ps[1:]:
        tc.remove(p)
    p0 = ps[0]
    for ch in list(p0):
        if ch.tag != qn('w:pPr'):
            p0.remove(ch)
    pPr = p0.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); p0.insert(0, pPr)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); pPr.append(jc)

    if code:
        full = code if code.startswith('MHZYY-') else 'MHZYY-' + code
        png = _make_qr_png(f'{host}/repair-fill?code={full}',
                           os.path.join(qr_dir, full + '.png'))
        if png:
            Paragraph(p0, doc).add_run().add_picture(png, width=Emu(int(QR_PX_CM * 360000)))
    _tight_para(cell, f'型号：{model}', 5.5)
    _tight_para(cell, f'编号：{code}', 5.5)
    _tight_para(cell, '设备故障请扫码', 5.5)
    _tight_para(cell, '填写维修记录', 5.5)
    return tbl


def build_docx(items, host: str, qr_dir: str) -> bytes:
    """生成设备卡片 Word，返回 bytes"""
    os.makedirs(qr_dir, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.2)
    sec.right_margin = Cm(3.2)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)

    for idx, item in enumerate(items):
        build_card_table(doc, item, qr_dir, host)
        # 卡片之间留空（每 PER_PAGE 张内留空行，页尾分页）
        if (idx + 1) % PER_PAGE == 0 and idx + 1 < len(items):
            doc.add_page_break()
        else:
            for _ in range(BLANK_LINES):
                doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', '_', name)[:80]
